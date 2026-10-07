#!/usr/bin/env python3
"""Builds Client-Finder-CRM.xlsx: the Google Sheet template buyers import (File > Import in Google Sheets).

Tabs:
  Leads       exactly the columns the workflows read/write (header row must stay as-is)
  Pipeline    live counts per status / priority (formulas)
  How to use  short instructions
Usage: python3 scripts/make-crm.py <output.xlsx>
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent


def node_json(expr):
    out = subprocess.run(["node", "-e", f"console.log(JSON.stringify({expr}))"], cwd=ROOT, check=True, capture_output=True, text=True)
    return json.loads(out.stdout)


COLUMNS = node_json("require('./src/lib/sheet').LEAD_COLUMNS")
STATUS = node_json("Object.values(require('./src/lib/sheet').STATUS)")

WIDTHS = {
    "Place ID": 14, "Business": 30, "Category": 18, "Phone": 16, "Email": 28, "Website": 30, "Address": 34,
    "Google Maps": 18, "Rating": 8, "Reviews": 9, "Score": 8, "Priority": 10, "Problems": 34, "Top Problem": 24,
    "Evidence": 60, "Email Subject": 34, "Email Body": 60, "DM Message": 44, "Call Opener": 44, "Status": 16,
    "Draft ID": 14, "Thread ID": 14, "Found On": 12, "Last Contacted": 14, "Follow-ups": 11, "Last Draft On": 13,
    "Niche": 16, "Search": 26, "AI Status": 26, "Notes": 40,
}

INK = "1F2937"
ACCENT = "4F46E5"
HEADER_FILL = PatternFill("solid", fgColor=INK)
THIN = Side(style="thin", color="E5E7EB")


def build(path):
    wb = Workbook()
    leads = wb.active
    leads.title = "Leads"
    leads.append(COLUMNS)
    for i, name in enumerate(COLUMNS, start=1):
        c = leads.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)
        leads.column_dimensions[get_column_letter(i)].width = WIDTHS.get(name, 16)
    leads.row_dimensions[1].height = 30
    leads.freeze_panes = "C2"
    last = get_column_letter(len(COLUMNS))
    leads.auto_filter.ref = f"A1:{last}1"

    col = {name: get_column_letter(i) for i, name in enumerate(COLUMNS, start=1)}
    # Status dropdown (the follow-up workflow understands every one of these).
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS) + '"', allow_blank=True, showDropDown=False)
    dv.error = "Pick a status from the list"
    dv.prompt = "New → Draft ready / Pitch ready → Contacted → Replied → Meeting booked → Won / Lost"
    leads.add_data_validation(dv)
    dv.add(f"{col['Status']}2:{col['Status']}5000")
    # Priority colours.
    pr = f"{col['Priority']}2:{col['Priority']}5000"
    leads.conditional_formatting.add(pr, CellIsRule(operator="equal", formula=['"Hot"'], fill=PatternFill("solid", fgColor="FEE2E2"), font=Font(color="B91C1C", bold=True)))
    leads.conditional_formatting.add(pr, CellIsRule(operator="equal", formula=['"Warm"'], fill=PatternFill("solid", fgColor="FEF3C7"), font=Font(color="92400E", bold=True)))
    leads.conditional_formatting.add(pr, CellIsRule(operator="equal", formula=['"Low"'], fill=PatternFill("solid", fgColor="F3F4F6"), font=Font(color="374151")))
    st = f"{col['Status']}2:{col['Status']}5000"
    for value, fill, font in [
        ("Replied", "DCFCE7", "166534"), ("Meeting booked", "DBEAFE", "1E40AF"), ("Won", "D1FAE5", "065F46"),
        ("Lost", "F3F4F6", "6B7280"), ("Do not contact", "F3F4F6", "6B7280"), ("Not interested", "F3F4F6", "6B7280"),
        ("Bounced", "FEE2E2", "991B1B"), ("Draft ready", "EEF2FF", "3730A3"),
    ]:
        leads.conditional_formatting.add(st, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=PatternFill("solid", fgColor=fill), font=Font(color=font)))
    leads.conditional_formatting.add(f"{col['Score']}2:{col['Score']}5000", FormulaRule(formula=[f"{col['Score']}2>=60"], font=Font(bold=True, color="B91C1C")))

    # Pipeline summary.
    pipe = wb.create_sheet("Pipeline")
    pipe["A1"] = "Client Finder pipeline"
    pipe["A1"].font = Font(bold=True, size=16, color=INK)
    pipe["A2"] = "Updates automatically from the Leads tab."
    pipe["A2"].font = Font(italic=True, color="6B7280")
    pipe.append([])
    pipe.append(["Status", "Leads"])
    for c in pipe[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
    for s in STATUS:
        pipe.append([s, f'=COUNTIF(Leads!{col["Status"]}:{col["Status"]},"{s}")'])
    pipe.append(["Total", f'=COUNTA(Leads!{col["Place ID"]}:{col["Place ID"]})-1'])
    pipe.cell(row=pipe.max_row, column=1).font = Font(bold=True)
    pipe.append([])
    start = pipe.max_row + 1
    pipe.append(["Priority", "Leads"])
    for c in pipe[start]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ACCENT)
    for p in ["Hot", "Warm", "Low", "None"]:
        pipe.append([p, f'=COUNTIF(Leads!{col["Priority"]}:{col["Priority"]},"{p}")'])
    pipe.append([])
    pipe.append(["Reply rate", f'=IFERROR((COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Replied")+COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Meeting booked")+COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Won"))/MAX(1,COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Contacted")+COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Follow-up drafted")+COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Replied")+COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Meeting booked")+COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Won")+COUNTIF(Leads!{col["Status"]}:{col["Status"]},"Lost")),0)'])
    pipe.cell(row=pipe.max_row, column=2).number_format = "0%"
    pipe.column_dimensions["A"].width = 22
    pipe.column_dimensions["B"].width = 12

    # How to use.
    how = wb.create_sheet("How to use")
    lines = [
        ("How to use this sheet", True),
        ("1. Keep the tab name \"Leads\" and its header row exactly as they are: the n8n workflows read and write these columns by name.", False),
        ("2. Copy this sheet's URL (the address bar, starting with https://docs.google.com/spreadsheets/d/) into the googleSheetUrl field of the ⚙️ Settings node.", False),
        ("3. Every run adds new businesses at the bottom. A business that is already in the sheet is never added twice, and an email address or website that is already here is never pitched again.", False),
        ("4. Statuses: New = no pitch (score too low, or already pitched under another listing) · Pitch ready = pitch written, no email found or the draft failed (call/DM) · Draft ready = a Gmail draft is waiting for you.", False),
        ("   After you send: Contacted · Follow-up drafted (Pro) · Replied · Meeting booked · Won · Lost · Do not contact · Not interested · Bounced (Pro sets it when an email bounces).", False),
        ("5. Pro follow-ups read Gmail to see what you actually sent and who replied (from any address in the thread). You can also set Status and Last Contacted by hand.", False),
        ("6. Set Status to \"Do not contact\" (or \"Not interested\") for anyone who asks you to stop. That email address is never emailed again by any workflow, and the Client Finder never pitches that address or website again, even under another listing.", False),
        ("7. Add your own columns to the right of Notes if you like: the workflows ignore them.", False),
        ("8. Business details come from Google Maps (source: Google Maps). Google's Maps Platform Terms limit storing them long term, so treat this tab as a working list and delete leads you don't pursue.", False),
    ]
    for text, bold in lines:
        how.append([text])
        how.cell(row=how.max_row, column=1).font = Font(bold=bold, size=14 if bold else 11, color=INK)
        how.cell(row=how.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    how.column_dimensions["A"].width = 120
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else ROOT / "products" / "Client-Finder-CRM.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    build(out)
    print(f"wrote {out}")
